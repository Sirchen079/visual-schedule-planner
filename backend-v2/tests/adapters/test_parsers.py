from zhishi.adapters.parsers import parse_file


def test_txt(tmp_path):
    path = tmp_path / 'sample.txt'
    path.write_text('示例日程', encoding='utf-8')
    doc = parse_file(path)
    assert doc.kind == 'text' and '示例日程' in doc.text and doc.tables == []


def test_csv(tmp_path):
    path = tmp_path / 'sample.csv'
    path.write_text('节次,星期一\n2,课程A[连续周1-16周]', encoding='utf-8')
    doc = parse_file(path)
    assert doc.kind == 'csv' and doc.tables[0][0] == ['节次', '星期一']


def test_docx_table(tmp_path):
    from docx import Document
    source = Document()
    table = source.add_table(rows=2, cols=3)
    rows = [['', '节次', '星期一'], ['', '2', '课程A[连续周1-16周]']]
    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    path = tmp_path / 'schedule.docx'
    source.save(path)
    doc = parse_file(path)
    assert doc.kind == 'docx' and doc.tables[0] == rows
    assert any('连续周' in block['text'] for block in doc.blocks)


def test_pdf_table(tmp_path):
    # A synthetic PDF keeps the parser test independent of external files.
    stream = (b'50 600 200 100 re S 50 650 m 250 650 l S 150 600 m 150 700 l S '
              b'BT /F1 12 Tf 60 670 Td (Item) Tj 100 0 Td (Time) Tj ET '
              b'BT /F1 12 Tf 60 620 Td (Meeting) Tj 100 0 Td (09:00) Tj ET')
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream',
    ]
    body = bytearray(b'%PDF-1.4\n')
    offsets = []
    for number, value in enumerate(objects, 1):
        offsets.append(len(body))
        body.extend(f'{number} 0 obj\n'.encode() + value + b'\nendobj\n')
    xref = len(body)
    body.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets:
        body.extend(f'{offset:010d} 00000 n \n'.encode())
    body.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    path = tmp_path / 'table.pdf'
    path.write_bytes(body)
    doc = parse_file(path)
    assert doc.kind == 'pdf'
    assert doc.tables[0] == [['Item', 'Time'], ['Meeting', '09:00']]
    assert 'Meeting' in doc.text


def test_unsupported_doc(tmp_path):
    assert parse_file(tmp_path / 'sample.doc').kind == 'unsupported'


def test_image_needs_vision(tmp_path):
    assert parse_file(tmp_path / 'sample.png').kind == 'image'
