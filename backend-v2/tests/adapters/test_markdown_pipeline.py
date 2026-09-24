"""Markdown 解析管道与扫描页 OCR 编排的针对性测试。

覆盖：docx 标题层级 → # 分节（blocks 带章节路径）、xlsx 合并单元格填充、
pptx 幻灯片/备注、PDF 扫描页检测、OCR 任务替换占位节并回填缓存
（成功/失败/kick 调度）、OCR 设置路由的 keyring 存取、sidecar 生命周期。"""
import asyncio
import json
import threading

import pytest

from zhishi.adapters import parsers
from zhishi.agent import ocr
from zhishi.agent.ocr import replace_page_section
from zhishi.domain.library import service


# ---- docx/xlsx/pptx → Markdown ----

def _write_docx(tmp_path, build):
    from docx import Document
    source = Document()
    build(source)
    path = tmp_path / 'doc.docx'
    source.save(path)
    return path


def test_docx_headings_become_markdown_sections_with_paths(tmp_path):
    def build(source):
        source.add_heading('第一章 总则', level=1)
        source.add_paragraph('第一章内容。')
        source.add_heading('数据来源', level=2)
        source.add_paragraph('数据来自问卷。')
        source.add_heading('附录', level=1)
        source.add_paragraph('附录内容。')
    doc = parsers.parse_file(_write_docx(tmp_path, build))
    assert '# 第一章 总则' in doc.markdown and '## 数据来源' in doc.markdown
    locations = [b['location'] for b in doc.blocks]
    # 章节路径分节：二级节带父级路径，read_material 引用因此可读
    assert any('第一章 总则 · 数据来源' == loc for loc in locations)
    assert any(loc == '附录' for loc in locations)


def test_docx_tables_preview_and_blocks(tmp_path):
    def build(source):
        source.add_heading('名单', level=1)
        table = source.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '姓名'
        table.cell(0, 1).text = '节次'
        table.cell(1, 0).text = '张三'
        table.cell(1, 1).text = '2'
    doc = parsers.parse_file(_write_docx(tmp_path, build))
    assert doc.tables == [[['姓名', '节次'], ['张三', '2']]]
    assert any('| 张三 | 2 |' in b['text'] for b in doc.blocks)


def test_xlsx_merged_cells_filled_and_sheet_sections(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    sheet = wb.active
    sheet.title = '期中'
    sheet['A1'] = '年级'
    sheet.merge_cells('A1:B2')
    sheet['C1'] = '高一'
    sheet.append(['一', '二', '三'])
    path = tmp_path / 'grades.xlsx'
    wb.save(path)
    doc = parsers.parse_file(path)
    assert '## 工作表「期中」' in doc.markdown
    assert doc.markdown.count('年级') >= 3   # 合并区域用左上值填充
    assert any(b['location'] == '工作表「期中」' for b in doc.blocks)


def test_pptx_slides_notes(tmp_path):
    from pptx import Presentation
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[5])
    slide.shapes.title.text = '课程导论'
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = '要点一'
    slide.notes_slide.notes_text_frame.text = '记得举例'
    path = tmp_path / 'deck.pptx'
    source.save(path)
    doc = parsers.parse_file(path)
    assert doc.kind == 'pptx'
    assert '## 幻灯片 1：课程导论' in doc.markdown
    assert '要点一' in doc.markdown and '**备注**：记得举例' in doc.markdown


def test_markdown_file_keeps_text_kind_and_headings_split_blocks(tmp_path):
    path = tmp_path / 'notes.md'
    path.write_text('# 标题甲\n内容甲\n\n# 标题乙\n内容乙', encoding='utf-8')
    doc = parsers.parse_file(path)
    assert doc.kind == 'text'
    assert any('标题乙' in b['location'] for b in doc.blocks)


# ---- PDF 扫描页检测 ----

def _synthetic_pdf(path, page_streams: list[bytes]) -> None:
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'']
    kids = []
    for number, stream in enumerate(page_streams, 1):
        page_id, content_id = len(objects) + 1, len(objects) + 2
        kids.append(f'{page_id} 0 R')
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] '
                       f'/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 '
                       f'/BaseFont /Helvetica >> >> >> /Contents {content_id} 0 R >>'.encode())
        objects.append(f'<< /Length {len(stream)} >>\nstream\n'.encode() + stream + b'\nendstream')
    objects[1] = f'<< /Type /Pages /Kids [{" ".join(kids)}] /Count {len(page_streams)} >>'.encode()
    pdf, offsets = b'%PDF-1.4\n', [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f'{number} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref = len(pdf)
    pdf += f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode()
    pdf += b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets[1:])
    pdf += f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
    path.write_bytes(pdf)


def test_pdf_marks_drawing_only_page_as_scan_placeholder(tmp_path):
    text_page = b'BT /F1 12 Tf 50 750 Td (PAGE 1 text) Tj ET'
    scan_page = b'50 50 500 700 re S'   # 只有矢量笔画没有文本：判扫描页
    blank_page = b''
    path = tmp_path / 'mixed.pdf'
    _synthetic_pdf(path, [text_page, scan_page, blank_page])
    doc = parsers.parse_file(path)
    assert doc.ocr_pages == [2]
    assert '## 第 2 页（扫描页）' in doc.markdown and 'PAGE 1 text' in doc.markdown
    assert any('扫描页' in w for w in doc.warnings)


# ---- OCR 占位节替换 ----

def test_replace_page_section_swaps_placeholder_and_appends_when_missing():
    md = ('## 第 1 页\n\n正文\n\n## 第 2 页（扫描页）\n\n（此页为扫描图像，等待 OCR 识别）\n\n'
          '## 第 3 页\n\n尾页')
    out = replace_page_section(md, 2, '识别出的内容')
    assert '识别出的内容' in out and '等待 OCR 识别' not in out
    assert '## 第 3 页' in out and '## 第 1 页' in out
    appended = replace_page_section('## 第 1 页\n\n正文', 2, '后补内容')
    assert appended.endswith('后补内容')


# ---- OCR 后台任务编排 ----

@pytest.fixture
def scan_file(db, tmp_path, monkeypatch):
    """含一个扫描页占位的已解析 PDF 文件行 + sidecar，模拟 ensure_parsed 后的 pending 态。"""
    from zhishi.domain.models import LibraryFile
    storage_root = tmp_path / 'attachments'
    storage_root.mkdir()
    source = storage_root / 'scan-source.pdf'   # 原文件本体不参与 OCR 测试（渲染被 mock）
    source.write_bytes(b'%PDF-1.4 fake')
    rel = source.relative_to(tmp_path).as_posix()
    row = LibraryFile(original_name='scan.pdf', storage_path=rel, size=4,
                      resource_type='file', parse_status='parsed', md_status='pending')
    db.add(row)
    db.commit()
    db.refresh(row)
    markdown = '## 第 1 页\n\n文字页\n\n## 第 2 页（扫描页）\n\n（此页为扫描图像，等待 OCR 识别）'
    service.write_markdown(storage_root, rel, markdown)
    doc = parsers.doc_from_markdown('pdf', markdown, warnings=['扫描页'])
    doc.ocr_pages = [2]
    row.parse_status, row.extracted_text = 'parsed', doc.to_json()
    db.commit()
    monkeypatch.setattr(ocr, '_render_page', lambda source, page, resolution=120: b'fake-png')
    return row, storage_root


async def test_ocr_convert_replaces_placeholder_and_finalizes(db, scan_file):
    row, storage_root = scan_file

    async def fake_ocr_image(model, png):
        assert png == b'fake-png'
        return 'OCR 识别的表格内容'
    monkeypatch_ocr_image(fake_ocr_image)
    factory = ocr._session_factory(db)
    await ocr._convert(factory, storage_root, row.id, {'base_url': 'http://x', 'model': 'm', 'api_key': 'k'})

    db.expire_all()
    fresh = db.get(type(row), row.id)
    assert fresh.md_status == 'done'
    assert 'OCR 识别的表格内容' in fresh.extracted_text
    assert json.loads(fresh.extracted_text)['ocr_pages'] == []   # 回填后不再有待办页
    md = service.markdown_path(storage_root, fresh.storage_path).read_text(encoding='utf-8')
    assert '## 第 2 页（扫描页 · OCR）' in md and '等待 OCR 识别' not in md


async def test_ocr_convert_marks_failed_on_page_error(db, scan_file, monkeypatch):
    row, storage_root = scan_file

    async def failing(model, png):
        raise RuntimeError('模型不可用')
    monkeypatch_ocr_image(failing)
    await ocr._convert(ocr._session_factory(db), storage_root, row.id,
                       {'base_url': 'http://x', 'model': 'm', 'api_key': 'k'})
    db.expire_all()
    fresh = db.get(type(row), row.id)
    assert fresh.md_status == 'failed'
    assert '等待 OCR 识别' in fresh.extracted_text   # 占位保留，可重试


def monkeypatch_ocr_image(fn):
    ocr.ocr_image = fn   # 模块级替身：_convert 内部按名字调用


async def test_kick_schedules_via_main_loop_and_dedupes(db, scan_file, monkeypatch):
    from zhishi.domain import settingsvc
    row, storage_root = scan_file
    settingsvc.set_setting(db, ocr.CONFIG_KEY, json.dumps(
        {'base_url': 'http://x', 'model': 'm'}))
    monkeypatch.setattr('zhishi.infra.secrets.load_api_key', lambda name: 'key')
    monkeypatch.setattr(ocr, '_main_loop', asyncio.get_running_loop())
    called = []
    real_convert = ocr._convert

    async def spy_convert(*args, **kwargs):
        called.append(args[2])
        await real_convert(*args, **kwargs)
    monkeypatch.setattr(ocr, '_convert', spy_convert)

    ocr.kick(db, row.id, storage_root)
    ocr.kick(db, row.id, storage_root)   # in-flight 去重：只投一次
    for _ in range(50):
        if called:
            break
        await asyncio.sleep(0.02)
    assert called == [row.id]
    ocr._inflight.discard(row.id)


def test_register_injects_service_hook():
    previous = service.ocr_kicker
    try:
        ocr.register()
        assert service.ocr_kicker is ocr.kick
    finally:
        service.ocr_kicker = previous


async def test_resume_pending_spawns_for_stuck_files(db, scan_file, monkeypatch):
    from zhishi.domain import settingsvc
    row, storage_root = scan_file
    settingsvc.set_setting(db, ocr.CONFIG_KEY, json.dumps({'base_url': 'http://x', 'model': 'm'}))
    monkeypatch.setattr('zhishi.infra.secrets.load_api_key', lambda name: 'key')
    monkeypatch.setattr(ocr, '_main_loop', asyncio.get_running_loop())
    resumed = ocr.resume_pending(ocr._session_factory(db), storage_root)
    assert resumed == 1
    ocr._inflight.discard(row.id)
    # 未配置 OCR 时不投递
    ocr._inflight.discard(row.id)
    settingsvc.set_setting(db, ocr.CONFIG_KEY, '')
    assert ocr.resume_pending(ocr._session_factory(db), storage_root) == 0


# ---- ensure_parsed 集成：sidecar / md_status / 缓存 ----

def test_ensure_parsed_writes_sidecar_and_marks_done(db, tmp_path):
    from docx import Document
    from zhishi.domain.models import LibraryFile
    storage_root = tmp_path / 'attachments'
    storage_root.mkdir()
    source = Document()
    source.add_heading('章节', level=1)
    source.add_paragraph('正文内容')
    target = storage_root / 'a.docx'
    source.save(target)
    row = LibraryFile(original_name='a.docx', storage_path='attachments/a.docx',
                      size=target.stat().st_size, resource_type='file', parse_status='pending')
    db.add(row)
    db.commit()
    db.refresh(row)
    doc = service.ensure_parsed(db, row, storage_root=storage_root)
    assert doc.kind == 'docx' and row.md_status == 'done'
    sidecar = service.markdown_path(storage_root, row.storage_path)
    assert sidecar.is_file() and '# 章节' in sidecar.read_text(encoding='utf-8')
    # 二次读取走缓存，sidecar 不被破坏
    service.ensure_parsed(db, row, storage_root=storage_root)
    assert sidecar.is_file()


def test_purge_removes_markdown_sidecar(db, tmp_path):
    from zhishi.domain.models import LibraryFile
    storage_root = tmp_path / 'attachments'
    storage_root.mkdir()
    target = storage_root / 'gone.txt'
    target.write_text('内容', encoding='utf-8')
    row = LibraryFile(original_name='gone.txt', storage_path='attachments/gone.txt',
                      size=6, resource_type='file', parse_status='parsed')
    db.add(row)
    db.commit()
    service.write_markdown(storage_root, row.storage_path, '# 副本')
    sidecar = service.markdown_path(storage_root, row.storage_path)
    assert sidecar.is_file()
    service.purge(db, row.id, storage_root=storage_root)
    assert not target.exists() and not sidecar.exists()


# ---- OCR 设置路由 ----

def test_ocr_settings_roundtrip(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from zhishi.server.app import create_app
    stored = {}
    monkeypatch.setattr('zhishi.infra.secrets.store_api_key',
                        lambda name, value: stored.__setitem__(name, value))
    monkeypatch.setattr('zhishi.infra.secrets.load_api_key',
                        lambda name: stored.get(name))
    monkeypatch.setattr('zhishi.infra.secrets.delete_api_key',
                        lambda name: stored.pop(name, None))
    with TestClient(create_app(data_dir=tmp_path)) as client:
        assert client.get('/api/settings/ocr').json() == {
            'base_url': '', 'model': '', 'has_api_key': False}
        body = {'base_url': 'https://api.siliconflow.cn/v1',
                'model': 'deepseek-ai/DeepSeek-OCR', 'api_key': 'sk-test'}
        assert client.put('/api/settings/ocr', json=body).json()['has_api_key'] is True
        got = client.get('/api/settings/ocr').json()
        assert got['model'] == 'deepseek-ai/DeepSeek-OCR' and got['has_api_key'] is True
        assert 'api_key' not in got   # 密钥永不回显
        # 通用设置接口不泄漏 key（key 只在 keyring）
        assert 'sk-test' not in json.dumps(client.get('/api/settings').json())
        assert client.put('/api/settings/ocr', json={'base_url': 'https://x', 'model': 'm',
                                                     'api_key': ''}).json()['has_api_key'] is False


# ---- 重新解析端点 ----

def test_reparse_endpoint_rebuilds_and_reports_md_status(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from zhishi.server.app import create_app
    with TestClient(create_app(data_dir=tmp_path)) as client:
        upload = client.post('/api/files', files={
            'file': ('doc.md', '# 标题\n第一段内容'.encode('utf-8'), 'text/markdown')})
        assert upload.status_code == 201
        file_id = upload.json()['id']
        detail = client.get(f'/api/files/{file_id}').json()
        assert detail['md_status'] in ('none', 'done')
        again = client.post(f'/api/files/{file_id}/reparse')
        assert again.status_code == 200
        assert again.json()['parse_status'] == 'parsed'
        assert again.json()['md_status'] == 'done'
        root = client.app.state.storage_root
        sidecar = root.parent / (again.json()['storage_path'] + '.md')
        assert sidecar.is_file() and '# 标题' in sidecar.read_text(encoding='utf-8')
